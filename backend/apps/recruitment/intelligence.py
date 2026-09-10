import hashlib
import logging
import re
import unicodedata
from datetime import date
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from django.db import transaction
from django.db.models import Prefetch

from .choices import ApplicationDocumentType
from .embeddings import EmbeddingError, embedding_model_identifier, semantic_score as calculate_semantic_score
from .models import Application, ApplicationDocument, ApplicationMatch, CVAnalysis, Offer, TrainingRecommendation
from .representations import (
    build_candidate_representation,
    build_offer_representation,
    candidate_matching_fingerprint,
    offer_matching_fingerprint,
    text_hash,
)

EXTRACTOR_VERSION = "structured-v7"
MATCHING_VERSION = "hybrid-v1"
MIN_USABLE_PDF_CHARACTERS = 60
SECTION_NAMES = {
    "skills": ("compétences", "competences", "skills", "technologies", "technical skills", "compétences techniques", "compétences professionnelles", "competences professionnelles", "compétences personnelles", "competences personnelles", "professional skills", "soft skills"),
    "experience": ("expériences", "experiences", "expérience", "experience", "expérience professionnelle", "expérience professionel", "experience professionel", "expériences professionnelles", "work experience", "professional experience", "employment", "parcours professionnel"),
    "education": ("formations", "formation", "éducation", "education", "diplômes", "diplomes", "academic background", "parcours académique", "parcours academique", "études", "etudes"),
    "languages": ("langues", "languages"),
    "certifications": ("certifications", "certificats", "certificates"),
    "interests": ("centres d’intérêt", "centres d'interet", "centres dinteret", "intérêts", "interets", "loisirs", "interests", "hobbies"),
    "availability": ("disponibilités", "disponibilites", "disponibilité", "disponibilite", "conditions de travail", "préférences de travail", "preferences de travail", "work preferences", "availability"),
    "projects": ("projets", "projet", "projets académiques", "projets academiques", "projects", "academic projects"),
    "profile": ("profil", "profile", "à propos", "a propos", "about me", "summary"),
    "engagements": ("engagements", "vie associative", "volunteering", "bénévolat", "benevolat"),
    "contact": ("contact", "coordonnées", "coordonnees", "contact details"),
    "references": ("références", "references"),
}
DATE_RANGE = re.compile(
    r"(?P<start>(?:0?[1-9]|1[0-2])[/.-]\d{4}|\d{4})\s*(?:[-–—à]|to)\s*"
    r"(?P<end>(?:0?[1-9]|1[0-2])[/.-]\d{4}|\d{4}|présent|present|actuel|current)", re.I
)
EMAIL = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PHONE_CANDIDATE = re.compile(
    r"(?<!\d)(?:(?:\+|00)?212[\s.()/-]*[67](?:[\s.()/-]*\d){8}|0[67](?:[\s.()/-]*\d){8})(?!\d)"
)
MONTH_DATE_RANGE = re.compile(
    r"(?P<start>(?:janvier|février|fevrier|mars|avril|mai|juin|juillet|août|aout|septembre|octobre|novembre|décembre|decembre)\s+\d{4}|\d{4})"
    r"\s*(?:[-–—à]|to)\s*(?P<end>(?:janvier|février|fevrier|mars|avril|mai|juin|juillet|août|aout|septembre|octobre|novembre|décembre|decembre)\s+\d{4}|\d{4}|présent|present|actuel|current)", re.I
)
MONTH_ONLY_RANGE = re.compile(
    r"(?P<start>janvier|février|fevrier|mars|avril|mai|juin|juillet|août|aout|septembre|octobre|novembre|décembre|decembre)"
    r"\s*(?:[-–—à]|to)\s*(?P<end>(?:janvier|février|fevrier|mars|avril|mai|juin|juillet|août|aout|septembre|octobre|novembre|décembre|decembre)\s+\d{4}|présent|present)", re.I
)

INSTITUTION_PATTERN = re.compile(
    r"\b(?:universit[ée]|university|[ée]cole|school|institut|institute|facult[ée]|academy|"
    r"groupe\s+scolaire|emsi|encg|ensa|ensam|est|fst|ofppt|ista|iscae|uir|um6p|hassan\s+ii)\b",
    re.I,
)
DEGREE_PATTERN = re.compile(
    r"\b(?:bac(?:calaur[ée]at)?|licence|bachelor|master|dipl[oô]me|doctorat|ph\.?d|"
    r"cycle\s+(?:d['’]\s*)?ing[ée]nierie|cycle\s+ing[ée]nieur|ing[ée]nieur|technicien|deug|dut|bts|"
    r"cycle\s+pr[ée]paratoire|\d(?:re|ère|e|ème)?\s+ann[ée]e)\b",
    re.I,
)
JOB_PATTERN = re.compile(
    r"\b(?:vendeu(?:r|se)|d[ée]veloppeu(?:r|se)|developer|engineer|ing[ée]nieur(?:e)?|consultant(?:e)?|"
    r"stagiaire|manager|responsable|assistant(?:e)?|comptable|commercial(?:e)?|"
    r"technicien(?:ne)?|analyste|designer|chef\s+de\s+projet|serveu(?:r|se)|"
    r"caissi(?:er|[èe]re)|enseignant(?:e)?|formateur|administrat(?:eur|rice))\b",
    re.I,
)
COMPANY_PATTERN = re.compile(
    r"\b(?:sarl|sa|sas|ltd|llc|inc|groupe|group|company|soci[ée]t[ée]|entreprise|cabinet)\b",
    re.I,
)
CITY_PATTERN = re.compile(
    r"\b(?:casablanca|rabat|marrakech|tanger|f[èe]s|agadir|mekn[èe]s|k[ée]nitra|"
    r"t[ée]touan|oujda|el\s+jadida|safi|mohammedia|maroc|morocco)\b",
    re.I,
)
LANGUAGE_PATTERN = re.compile(
    r"^(?:fran[çc]ais|french|anglais|english|arabe|arabic|espagnol|spanish|allemand|german|"
    r"italien|italian|amazigh|tamazight)(?:\s*[:\-–—]\s*.+)?$",
    re.I,
)
WORK_CONDITION_PATTERN = re.compile(
    r"^(?:week[ -]?end|présentiel|presentiel|télétravail|teletravail|hybride|remote|"
    r"horaires? décalés?|travail (?:le soir|de nuit)|soir|nuit|temps plein|temps partiel|"
    r"disponible|disponibilité|mobilité|immediate|immédiate)$",
    re.I,
)
TECH_SKILL_PATTERN = re.compile(
    r"\b(?:python|django|flask|fastapi|java|spring|javascript|typescript|angular|react(?:\.js)?|"
    r"vue(?:\.js)?|node(?:\.js)?|php|laravel|symfony|c\+\+|c#|\.net|sql|postgresql|mysql|"
    r"mongodb|docker|kubernetes|git|linux|aws|azure|gcp|power\s*bi|excel|tensorflow|"
    r"machine\s+learning|rest\s*api|drf|postgres|python3|js)\b",
    re.I,
)

logger = logging.getLogger(__name__)

SKILL_ALIASES = {
    "python3": "Python", "python": "Python", "django": "Django",
    "drf": "Django REST Framework", "django rest framework": "Django REST Framework",
    "js": "JavaScript", "javascript": "JavaScript", "typescript": "TypeScript",
    "postgres": "PostgreSQL", "postgresql": "PostgreSQL", "mysql": "MySQL",
    "rest api": "REST API", "restful api": "REST API", "api rest": "REST API",
    "ml": "Machine Learning", "machine learning": "Machine Learning",
    "power bi": "Power BI", "tensorflow": "TensorFlow", "git": "Git",
    "docker": "Docker", "angular": "Angular", "java": "Java", "sql": "SQL",
}


class CVExtractionError(ValueError):
    pass


class MatchingCalculationError(ValueError):
    """Raised when no honest application score can be calculated."""


@dataclass
class ExtractedDocument:
    text: str
    method: str
    warnings: list[str]


def _extract_pdf(file_bytes: bytes) -> ExtractedDocument:
    import pymupdf

    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            text = "\n".join(_ordered_pdf_page_text(page) for page in document)
    except (pymupdf.FileDataError, RuntimeError, ValueError) as exc:
        raise CVExtractionError("Le CV PDF est invalide ou corrompu.") from exc

    if _usable_character_count(text) >= MIN_USABLE_PDF_CHARACTERS:
        return ExtractedDocument(text, "PDF_TEXT", [])
    try:
        ocr_text = _ocr_pdf(file_bytes)
    except CVExtractionError as exc:
        raise CVExtractionError(
            "Ce PDF ne contient pas de texte exploitable en extraction normale et semble scanné. "
            f"{exc}"
        ) from exc
    if _usable_character_count(ocr_text) < MIN_USABLE_PDF_CHARACTERS:
        raise CVExtractionError(
            "Ce PDF ne contient pas de texte exploitable, même après exécution de l’OCR."
        )
    return ExtractedDocument(
        ocr_text,
        "PDF_OCR",
        ["Le PDF ne contenait presque aucun texte : une extraction OCR de secours a été utilisée."],
    )


def _usable_character_count(text: str) -> int:
    return sum(character.isalnum() for character in text)


def _ocr_pdf(file_bytes: bytes) -> str:
    """Best-effort OCR fallback. Imports remain optional by design."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise CVExtractionError(
            "Le composant OCR optionnel n’est pas installé ; fournissez un PDF textuel ou un DOCX."
        ) from exc
    try:
        texts = []
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            for page in document:
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                image = Image.open(BytesIO(pixmap.tobytes("png")))
                texts.append(pytesseract.image_to_string(image, lang="fra+eng"))
        return "\n".join(texts)
    except Exception as exc:
        raise CVExtractionError(
            "L’OCR optionnel est indisponible ou n’a pas pu lire ce document."
        ) from exc


def _ordered_pdf_page_text(page) -> str:
    """Rebuild a stable reading order from positioned text blocks.

    PyMuPDF's plain-text order follows the PDF content stream, which can
    alternate unrelated columns. When both page halves contain several
    blocks, read each column top-to-bottom; otherwise preserve visual order.
    """
    blocks = [block for block in page.get_text("blocks") if block[6] == 0 and block[4].strip()]
    if not blocks:
        return ""
    starts = sorted(set(block[0] for block in blocks))
    gaps = [(right - left, left, right) for left, right in zip(starts, starts[1:])]
    largest_gap = max(gaps, default=(0, 0, 0))
    if largest_gap[0] >= page.rect.width * 0.12:
        split = (largest_gap[1] + largest_gap[2]) / 2
        left = [block for block in blocks if block[0] < split]
        right = [block for block in blocks if block[0] >= split]
    else:
        left, right = [], []
    if len(left) >= 3 and len(right) >= 3:
        ordered = sorted(left, key=lambda block: (block[1], block[0])) + sorted(
            right, key=lambda block: (block[1], block[0])
        )
    else:
        ordered = sorted(blocks, key=lambda block: (round(block[1] / 4), block[0]))
    return "\n".join(block[4].strip() for block in ordered)


def _extract_docx(file_bytes: bytes) -> ExtractedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise CVExtractionError("Le CV DOCX est invalide ou corrompu.") from exc
    parts = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            value = Paragraph(child, document).text.strip()
            if value:
                parts.append(value)
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    parts.extend(value.strip() for value in cell.text.splitlines() if value.strip())
    return ExtractedDocument("\n".join(parts), "DOCX", [])


def extract_document(uploaded_file, filename: str | None = None) -> tuple[ExtractedDocument, bytes]:
    uploaded_file.seek(0)
    raw = uploaded_file.read()
    uploaded_file.seek(0)
    suffix = Path(filename or getattr(uploaded_file, "name", "")).suffix.lower()
    if suffix == ".pdf":
        result = _extract_pdf(raw)
    elif suffix == ".docx":
        result = _extract_docx(raw)
    elif suffix == ".txt":  # retained for deterministic unit tests only
        result = ExtractedDocument(raw.decode("utf-8", errors="ignore"), "TEXT", [])
    else:
        raise CVExtractionError("Format de CV non pris en charge. Utilisez PDF ou DOCX.")
    if not result.text.strip():
        raise CVExtractionError("Le CV ne contient aucun texte exploitable.")
    return result, raw


def _repair_character_spaced_line(line: str) -> str:
    """Repair PDFs whose embedded font maps emit one token per character.

    Original word boundaries are represented by repeated spaces, while letters
    inside a word have one space. Normal lines are returned unchanged.
    """
    stripped = line.strip()
    tokens = stripped.split()
    if len(tokens) < 4:
        return stripped
    single = sum(len(token) == 1 for token in tokens)
    if single / len(tokens) < 0.68:
        return stripped
    words = []
    for group in re.split(r"\s{2,}", stripped):
        compact = re.sub(r"(?<=\S)\s(?=\S)", "", group)
        if compact:
            words.append(compact)
    return " ".join(words)


def normalize_extracted_text(text: str) -> str:
    repaired = "\n".join(_repair_character_spaced_line(line) for line in text.splitlines())
    # Multi-column PDFs sometimes concatenate the next uppercase heading to
    # the final value of the preceding section.
    return re.sub(r"(?<=[a-zà-ÿé])(?=[A-ZÉÈÀÙÇ]{5,}(?:\s|$))", "\n", repaired)


def _clean_lines(text: str) -> list[str]:
    text = normalize_extracted_text(text)
    lines = []
    for line in text.splitlines():
        line = re.sub(r"(?:\s*\|\s*){2,}", " ", line)
        line = re.sub(r"\s+", " ", line).strip(" •·|\t")
        if line:
            lines.append(line)
    return lines


def _normalized_heading(line: str) -> str:
    return re.sub(r"[^a-zà-ÿ ]", "", line.casefold()).strip()


def _heading_match(line: str) -> tuple[str, str]:
    """Return a section and optional inline value for an explicit heading."""
    stripped = line.strip()
    before, separator, after = stripped.partition(":")
    candidate = _normalized_heading(before if separator else stripped)
    for name, labels in SECTION_NAMES.items():
        if candidate in {_normalized_heading(label) for label in labels}:
            return name, after.strip() if separator else ""
    return "", ""


def detect_sections(lines: list[str]) -> dict[str, list[str]]:
    """Assign lines to sections without interpreting their field structure."""
    values = {name: [] for name in SECTION_NAMES}
    current = ""
    semantic_education = False
    for line in lines:
        found, remainder = _heading_match(line)
        if found:
            current = found
            semantic_education = False
            if remainder:
                values[current].append(remainder)
            continue

        if LANGUAGE_PATTERN.fullmatch(line.strip(" •")):
            values["languages"].append(line)
            semantic_education = False
            continue
        education_signal = bool(
            DEGREE_PATTERN.search(line)
            and current in {"", "education", "experience"}
            and len(line) <= 120
            and not re.search(r"\b(?:je|i am|i'm|étudiant(?:e)?|student)\b", line, re.I)
        )
        education_continuation = semantic_education and bool(
            INSTITUTION_PATTERN.search(line)
            or _date_match(line)
            or (line.isupper() and len(line.split()) <= 8)
        )
        if education_signal or education_continuation:
            values["education"].append(line)
            semantic_education = True
            continue
        semantic_education = False
        if current:
            values[current].append(line)
    return values


def _unique(values):
    result = []
    seen = set()
    for value in values:
        value = re.sub(r"(?:\s*\|\s*)+", " ", value).strip(" ,;•·-")
        value = re.sub(r"\s+", " ", value)
        key = _fold(value)
        if value and key not in seen:
            seen.add(key); result.append(value)
    return result


def _split_items(lines: list[str]) -> list[str]:
    return _unique(item for line in lines for item in re.split(r"[,;|•·]", line) if 1 < len(item.strip()) < 100)


def _skill_items(lines: list[str]) -> list[str]:
    values = []
    for line in lines:
        # Category labels such as "Frameworks:" describe the following values;
        # they are not themselves candidate skills.
        content = line.split(":", 1)[1] if ":" in line else line
        values.extend(re.split(r"[,;|•·]", content))
    return _unique(
        value for value in values
        if 1 < len(value.strip()) < 80
        and not WORK_CONDITION_PATTERN.fullmatch(value.strip())
        and not INSTITUTION_PATTERN.search(value)
        and not DEGREE_PATTERN.search(value)
        and not _heading_match(value)[0]
    )


def _normalize_skills(skills: list[str]) -> list[str]:
    folded = {_fold(skill): skill for skill in skills}
    if "react" in folded and "native" in folded:
        skills = [skill for skill in skills if _fold(skill) not in {"react", "native"}]
        skills.append("React Native")
    normalized = []
    for skill in skills:
        cleaned = re.sub(r"\s+", " ", skill).strip()
        normalized.append(SKILL_ALIASES.get(_fold(cleaned), cleaned))
    return _unique(normalized)


def _fold(value: str) -> str:
    return "".join(char for char in unicodedata.normalize("NFKD", value.casefold()) if not unicodedata.combining(char))


def _heading_kind(line: str) -> str:
    return _heading_match(line)[0]


def _section_indexes(lines: list[str]) -> dict[int, str]:
    result: dict[int, str] = {}
    current = ""
    for index, line in enumerate(lines):
        found = _heading_kind(line)
        if found:
            current = found; result[index] = "heading"
        elif current:
            result[index] = current
    return result


def _normalize_phone(value: str) -> str:
    digits = re.sub(r"\D", "", value)
    if digits.startswith("00212"):
        digits = digits[2:]
    if re.fullmatch(r"212[67]\d{8}", digits):
        return f"+{digits}"
    if re.fullmatch(r"0[67]\d{8}", digits):
        return f"+212{digits[1:]}"
    return ""


def _extract_phones(lines: list[str]) -> list[str]:
    return _unique(
        normalized for line in lines for match in PHONE_CANDIDATE.finditer(line)
        if (normalized := _normalize_phone(match.group(0)))
    )


def _plausible_name_line(line: str) -> bool:
    words = line.strip().split()
    if not 1 <= len(words) <= 5 or not 2 <= len(line) <= 80:
        return False
    if any(char.isdigit() for char in line) or EMAIL.search(line) or PHONE_CANDIDATE.search(line):
        return False
    if not all(re.fullmatch(r"[A-Za-zÀ-ÖØ-öø-ÿ'’-]+", word) for word in words):
        return False
    if JOB_PATTERN.search(line) or INSTITUTION_PATTERN.search(line) or DEGREE_PATTERN.search(line):
        return False
    if COMPANY_PATTERN.search(line) or CITY_PATTERN.fullmatch(line.strip()):
        return False
    return _fold(line) not in {
        "curriculum vitae", "cv", "profil", "profile", "contact", "competences", "skills",
        "experience", "experiences", "formation", "formations", "education", "langues",
        "languages", "certifications", "loisirs", "interets", "centres dinteret",
        "disponibilites", "conditions de travail", "references",
    }


def _name_score(index: int, line: str, lines: list[str], contact_indexes: set[int]) -> int:
    if not _plausible_name_line(line):
        return -100
    score = max(0, 34 - index * 3)
    if 2 <= len(line.split()) <= 4:
        score += 22
    if line.isupper() or line.istitle():
        score += 8
    distance = min((abs(index - contact) for contact in contact_indexes), default=99)
    if distance <= 2:
        score += 24 - distance * 5
    email_parts = {
        _fold(part) for email_line in lines for email in EMAIL.findall(email_line)
        for part in re.split(r"[._+-]", email.split("@", 1)[0]) if len(part) > 1
    }
    score += 8 * sum(_fold(word) in email_parts for word in line.split())
    email_locals = [_fold(email.split("@", 1)[0]) for email_line in lines for email in EMAIL.findall(email_line)]
    score += 12 * sum(any(_fold(word) in local for local in email_locals) for word in line.split())
    previous = lines[index - 1] if index else ""
    following = lines[index + 1] if index + 1 < len(lines) else ""
    if JOB_PATTERN.search(previous) or DATE_RANGE.search(following):
        score -= 45
    if INSTITUTION_PATTERN.search(previous) or INSTITUTION_PATTERN.search(following):
        score -= 30
    return score


def _detect_name(lines: list[str], section_indexes: dict[int, str]) -> tuple[str, int]:
    contact_indexes = {index for index, line in enumerate(lines) if EMAIL.search(line) or PHONE_CANDIDATE.search(line)}
    candidates: list[tuple[int, str, int]] = []
    for index, line in enumerate(lines[:60]):
        near_contact = any(abs(index - contact) <= 2 for contact in contact_indexes)
        email_related = any(
            _fold(word) in _fold(email.split("@", 1)[0])
            for word in line.split()
            for email_line in lines
            for email in EMAIL.findall(email_line)
            if len(word) > 2
        )
        # Multi-column extraction can place the visual header after a section.
        # In that case proximity to a validated contact may override section order.
        if (section_indexes.get(index) and not near_contact and not email_related) or re.search(r"https?://|linkedin|curriculum|vitae", line, re.I):
            continue
        score = _name_score(index, line, lines, contact_indexes)
        if score >= 35:
            candidates.append((score, line, index))
        if (
            len(line.split()) == 1
            and index + 1 < min(len(lines), 60)
            and (not section_indexes.get(index + 1) or near_contact)
            and len(lines[index + 1].split()) == 1
            and _plausible_name_line(line)
            and _plausible_name_line(lines[index + 1])
            and not _heading_kind(line)
            and not _heading_kind(lines[index + 1])
        ):
            combined = f"{line} {lines[index + 1]}"
            pair_score = _name_score(index, combined, lines, contact_indexes) + 10
            if pair_score >= 35:
                candidates.append((pair_score, combined, index))
    if not candidates:
        return "", 0
    score, value, _ = max(candidates, key=lambda item: (item[0], -item[2]))
    return value, min(100, score)


def _city_from_line(line: str) -> str:
    matches = list(CITY_PATTERN.finditer(line))
    if not matches:
        return ""
    return next((match.group(0) for match in matches if _fold(match.group(0)) not in {"maroc", "morocco"}), matches[0].group(0))


def _detect_location(lines: list[str]) -> tuple[str, int]:
    for line in lines[:25]:
        match = re.match(r"(?:adresse|address|localisation|location|ville|city)\s*[:|-]\s*(.+)", line, re.I)
        if match:
            value = match.group(1).strip()
            if not INSTITUTION_PATTERN.search(value) and not COMPANY_PATTERN.search(value):
                return value, 95
            city = _city_from_line(value)
            return (city, 65) if city else ("", 0)
    for line in lines[:25]:
        city = _city_from_line(line)
        if not city:
            continue
        if INSTITUTION_PATTERN.search(line) or COMPANY_PATTERN.search(line):
            return city, 55
        if len(line.split()) <= 5:
            return line.strip(), 75
    return "", 0


def _institution_and_city(value: str) -> tuple[str, str]:
    city = _city_from_line(value)
    institution = value
    if city:
        institution = re.sub(rf"\s*[-–—|,]\s*{re.escape(city)}(?:\s*[-–—|,]\s*(?:Maroc|Morocco))?\s*$", "", institution, flags=re.I)
    return institution.strip(" -–—|,"), city


def _validate_parsed_data(data: dict) -> dict:
    name = data.get("full_name", "").strip()
    conflicts = {_fold(value) for value in [*data.get("companies", []), *data.get("positions", [])] if value}
    if not _plausible_name_line(name) or _fold(name) in conflicts:
        name = ""
    data["full_name"] = name
    data["first_name"] = name.split()[0] if name else ""
    data["last_name"] = " ".join(name.split()[1:]) if name else ""
    data["phone"] = _normalize_phone(data.get("phone", ""))
    data["contact_details"]["phones"] = [data["phone"]] if data["phone"] else []
    location = data.get("location", "").strip()
    if INSTITUTION_PATTERN.search(location) or COMPANY_PATTERN.search(location):
        location = _city_from_line(location)
    data["location"] = location
    data["contact_details"]["location"] = location
    return data


def _date_match(value: str):
    return DATE_RANGE.search(value) or MONTH_DATE_RANGE.search(value) or MONTH_ONLY_RANGE.search(value)


def _without_date(value: str) -> str:
    value = DATE_RANGE.sub("", value)
    value = MONTH_DATE_RANGE.sub("", value)
    value = MONTH_ONLY_RANGE.sub("", value)
    return re.sub(r"(?:\s*\|\s*){2,}", " ", value).strip(" -–—|,")


def _normalize_experience_position(value: str, max_length: int = 255) -> str:
    """Keep an extracted position as a title, never as a responsibility paragraph."""
    value = re.sub(r"\s+", " ", value).strip(" \t\r\n-–—|,;:•")
    if not value:
        return ""
    # OCR/layout extraction can join the title to bullets or responsibility
    # sentences. The complete source remains available in ``description``.
    value = re.split(r"\s*(?:[•●▪◦]|\s[-–—]\s|[.;]\s+)\s*", value, maxsplit=1)[0].strip()
    if len(value) <= max_length:
        return value
    shortened = value[:max_length].rstrip()
    if " " in shortened:
        shortened = shortened.rsplit(" ", 1)[0]
    return shortened.rstrip(" -–—|,;:")


def _experience_records(lines: list[str]) -> list[dict]:
    records = []
    pending: list[str] = []
    groups: list[list[str]] = []
    for line in lines:
        cleaned = re.sub(r"(?:\s*\|\s*){2,}", " ", line).strip(" |")
        if not cleaned:
            continue
        pending.append(cleaned)
        if _date_match(cleaned):
            groups.append(pending); pending = []
        elif len(pending) >= 3 and JOB_PATTERN.search(pending[-1]) and JOB_PATTERN.search(pending[0]):
            groups.append(pending[:-1]); pending = [pending[-1]]
    if pending:
        groups.append(pending)

    for group in groups:
        date_line = next((line for line in group if _date_match(line)), "")
        date = _date_match(date_line) if date_line else None
        content = [_without_date(line) for line in group if _without_date(line)]
        if not content:
            continue
        chunks = [part.strip() for part in re.split(r"\s+\|\s+|\s+@\s+", " | ".join(content)) if part.strip()]
        position_index = next((i for i, value in enumerate(chunks) if JOB_PATTERN.search(value)), None)
        position = chunks[position_index] if position_index is not None else ""
        company = ""
        if position:
            combined = [part.strip() for part in re.split(r"\s+[-–—]\s+", position, maxsplit=1)]
            if len(combined) == 2 and JOB_PATTERN.search(combined[0]):
                position, company = combined
            if not company:
                company = next(
                    (
                        value for i, value in enumerate(chunks)
                        if i != position_index
                        and not CITY_PATTERN.fullmatch(value.strip())
                        and not DEGREE_PATTERN.search(value)
                        and not INSTITUTION_PATTERN.search(value)
                    ),
                    "",
                )
        position = _normalize_experience_position(position)
        description = "\n".join(_unique(content))
        record = {
            "position": position, "company": company,
            "start_date": date.group("start") if date else "",
            "end_date": date.group("end") if date else "",
            "duration": date_line.strip(" |") if date else "",
            "description": description,
        }
        records.append(record)
    merged = []
    for record in records:
        if merged and not record["start_date"] and not record["end_date"]:
            previous = merged[-1]
            previous["description"] = "\n".join(
                _unique([previous["description"], record["description"]])
            )
            if not previous["position"] and record["position"]:
                previous["position"] = record["position"]
            if not previous["company"] and record["company"]:
                previous["company"] = record["company"]
        else:
            merged.append(record)
    result, seen = [], set()
    for record in merged:
        if record["position"] or record["company"]:
            key = tuple(_fold(record[field]) for field in ("position", "company", "start_date", "end_date"))
        else:
            key = tuple(_fold(record[field]) for field in ("start_date", "end_date", "description"))
        if key not in seen:
            seen.add(key); result.append(record)
    return result


def _education_records(lines: list[str]) -> list[dict]:
    records = []
    current = None
    for line in lines:
        if EMAIL.search(line) or PHONE_CANDIDATE.search(line):
            continue
        date = _date_match(line)
        cleaned = _without_date(line)
        if not cleaned and date:
            if current and not current["institution"] and not current["start_date"] and not current["end_date"]:
                current["start_date"], current["end_date"] = date.group("start"), date.group("end")
            continue
        parts = [part.strip() for part in re.split(r"\s+\|\s+|\s+@\s+", cleaned) if part.strip()]
        for part in parts:
            if INSTITUTION_PATTERN.search(part):
                institution, _ = _institution_and_city(part)
                if current and not current["institution"]:
                    current["institution"] = institution
                else:
                    current = {"title": "", "institution": institution, "start_date": "", "end_date": ""}
                    records.append(current)
            elif DEGREE_PATTERN.search(part):
                if current and not current["title"]:
                    current["title"] = part
                else:
                    current = {"title": part, "institution": "", "start_date": "", "end_date": ""}
                    records.append(current)
            elif current and not current["title"]:
                current["title"] = part
            elif current and not current["institution"] and part.isupper() and len(part.split()) <= 8:
                current["title"] = f'{current["title"]} {part}'.strip()
            if date and current:
                if not current["start_date"] and not current["end_date"]:
                    current["start_date"], current["end_date"] = date.group("start"), date.group("end")
    return _unique_records(records, ("title", "institution", "start_date", "end_date"))


def _unique_records(records: list[dict], fields: tuple[str, ...]) -> list[dict]:
    result, seen = [], set()
    for record in records:
        key = tuple(_fold(str(record.get(field, ""))) for field in fields)
        if any(key) and key not in seen:
            seen.add(key); result.append(record)
    return result


def parse_cv_text(text: str) -> dict:
    lines = _clean_lines(text)
    normalized_text = "\n".join(lines)
    sections = detect_sections(lines)
    section_indexes = _section_indexes(lines)
    emails = _unique(EMAIL.findall(normalized_text))
    phones = _extract_phones(lines)
    name, name_confidence = _detect_name(lines, section_indexes)
    location, location_confidence = _detect_location(lines)

    experiences = _experience_records(sections["experience"])
    companies = _unique(item["company"] for item in experiences if item["company"])
    positions = _unique(item["position"] for item in experiences if item["position"])
    education = _education_records(sections["education"])
    certifications = _split_items(sections["certifications"])
    languages = _unique(
        item for item in _split_items(sections["languages"])
        if LANGUAGE_PATTERN.fullmatch(item.strip())
    )
    skills = _normalize_skills(
        [item for item in _skill_items(sections["skills"]) if not DATE_RANGE.search(item)]
    )
    if not skills:
        # Without an explicit section, only accept a conservative technology
        # vocabulary; never promote arbitrary neighbouring labels to skills.
        skills = _normalize_skills(
            sorted(_unique(match.group(0).casefold() for match in TECH_SKILL_PATTERN.finditer(normalized_text)))
        )
    data = {
        "first_name": name.split()[0] if name else "",
        "last_name": " ".join(name.split()[1:]) if name else "",
        "full_name": name,
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "location": location,
        "skills": skills,
        "experiences": experiences,
        "education": education,
        "diplomas": [item["title"] for item in education if item["title"]],
        "companies": _unique(companies),
        "positions": _unique(positions),
        "languages": languages,
        "certifications": certifications,
        "contact_details": {
            "emails": emails, "phones": phones, "location": location,
            "confidence": {"name": name_confidence, "email": 100 if emails else 0,
                           "phone": 100 if phones else 0, "location": location_confidence},
        },
    }
    return _validate_parsed_data(data)


def extract_cv_data(uploaded_file, filename: str | None = None, *, include_raw_text=False) -> dict:
    document, raw = extract_document(uploaded_file, filename)
    data = parse_cv_text(document.text)
    data.update({"source_sha256": hashlib.sha256(raw).hexdigest(), "extraction_method": document.method, "extraction_warnings": document.warnings, "extractor_version": EXTRACTOR_VERSION})
    if include_raw_text:
        data["raw_text"] = document.text
    return data


def cv_document_sha256(document: ApplicationDocument) -> str:
    """Hash the stored bytes of a CV without relying on names or timestamps."""
    with document.file.open("rb") as source:
        digest = hashlib.sha256()
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_cv(application: Application, *, force=False) -> CVAnalysis:
    document = application.documents.filter(document_type=ApplicationDocumentType.CV).order_by("-uploaded_at", "-id").first()
    if not document:
        raise CVExtractionError("Aucun CV n’est associé à cette candidature.")
    with document.file.open("rb") as source:
        data = extract_cv_data(source, document.original_name or document.file.name, include_raw_text=True)
    data.pop("first_name", None)
    data.pop("last_name", None)
    existing = getattr(application, "cv_analysis", None)
    if existing and existing.human_validated and not force:
        raise CVExtractionError("Les données ont déjà été corrigées et validées. Confirmez la réanalyse.")
    defaults = {**data, "human_validated": False, "validated_by": None, "validated_at": None}
    analysis, _ = CVAnalysis.objects.update_or_create(application=application, defaults=defaults)
    return analysis


def _skill_set(values) -> dict[str, str]:
    result = {}
    for value in values:
        canonical = _normalize_skills([value])
        if canonical:
            result[_fold(canonical[0])] = canonical[0]
    return result


def _offer_skills(offer: Offer) -> list[str]:
    return _normalize_skills(
        [item for item in re.split(r"[,;\n|]", offer.required_skills) if item.strip()]
    )


def _parse_month(value: str, *, end=False) -> int | None:
    folded = _fold(value.strip())
    if not folded:
        return None
    if folded in {"present", "actuel", "current", "aujourd'hui"}:
        today = date.today()
        return today.year * 12 + today.month
    numeric = re.fullmatch(r"(?:(0?[1-9]|1[0-2])[/.-])?(\d{4})", folded)
    if numeric:
        month = int(numeric.group(1) or (12 if end else 1))
        return int(numeric.group(2)) * 12 + month
    months = {
        "janvier": 1, "january": 1, "fevrier": 2, "february": 2,
        "mars": 3, "march": 3, "avril": 4, "april": 4, "mai": 5, "may": 5,
        "juin": 6, "june": 6, "juillet": 7, "july": 7, "aout": 8, "august": 8,
        "septembre": 9, "september": 9, "octobre": 10, "october": 10,
        "novembre": 11, "november": 11, "decembre": 12, "december": 12,
    }
    named = re.fullmatch(r"([a-z]+)\s+(\d{4})", folded)
    if named and named.group(1) in months:
        return int(named.group(2)) * 12 + months[named.group(1)]
    return None


def approximate_experience_years(experiences: list[dict]) -> float | None:
    """Conservatively merge dated employment intervals and return years."""
    intervals = []
    for experience in experiences:
        start = _parse_month(str(experience.get("start_date", "")))
        end = _parse_month(str(experience.get("end_date", "")), end=True)
        if start is not None and end is not None and end >= start:
            intervals.append((start, end))
    if not intervals:
        return None
    intervals.sort()
    merged = []
    for start, end in intervals:
        if merged and start <= merged[-1][1] + 1:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    months = sum(end - start + 1 for start, end in merged)
    return round(months / 12, 1)


EDUCATION_RANK = {
    "FIRST_YEAR": 1, "SECOND_YEAR": 2, "THIRD_YEAR": 3, "DUT": 3, "BTS": 3,
    "BACHELOR": 3, "FOURTH_YEAR": 4, "MASTER": 5, "FIFTH_YEAR": 5,
    "ENGINEERING": 5, "DOCTORATE": 6,
}


def _candidate_education_rank(application: Application, analysis: CVAnalysis) -> int | None:
    detected = " ".join(
        [*analysis.diplomas, *(item.get("title", "") for item in analysis.education)]
    )
    folded = _fold(detected)
    text_rank = None
    for pattern, rank in (
        (r"doctorat|ph\.?d", 6), (r"ingenieur|master", 5),
        (r"licence|bachelor", 3), (r"\bdut\b|\bbts\b", 3),
    ):
        if re.search(pattern, folded):
            text_rank = rank
            break
    profile_rank = EDUCATION_RANK.get(application.candidate_profile.study_level)
    return max((rank for rank in (text_rank, profile_rank) if rank is not None), default=None)


def _score_label(score: float) -> str:
    if score >= 90: return "Excellent match"
    if score >= 75: return "Strong match"
    if score >= 60: return "Moderate match"
    if score >= 40: return "Weak match"
    return "Very weak match"


def _candidate_summary(*, matched, missing, experience_years, education_score, score) -> str:
    sentences = []
    if matched:
        sentences.append(f"Le profil correspond notamment sur : {', '.join(matched[:6])}.")
    else:
        sentences.append("Aucune compétence requise n’a été identifiée avec suffisamment de certitude.")
    if experience_years is not None:
        sentences.append(f"Environ {experience_years:g} année(s) d’expérience datée ont été identifiées.")
    if education_score is not None:
        sentences.append("Le niveau d’études détecté répond au critère de l’offre." if education_score == 100 else "Le niveau d’études détecté ne permet pas de confirmer entièrement le critère de l’offre.")
    if missing:
        sentences.append(f"Principales compétences manquantes : {', '.join(missing[:6])}.")
    sentences.append(f"Correspondance globale : {_score_label(score)} ({score:.0f} %). Cette note est une aide à la décision et ne provoque aucun rejet automatique.")
    return " ".join(sentences)


def normalized_weighted_score(weighted: list[tuple[float, int]]) -> tuple[float, int]:
    """Normalize available criteria to 100 without treating missing ones as zero."""
    if not weighted:
        raise MatchingCalculationError(
            "Aucun critère exploitable n’est disponible pour calculer un score fiable."
        )
    total_weight = sum(weight for _, weight in weighted)
    if total_weight <= 0:
        raise MatchingCalculationError("La somme des poids disponibles doit être positive.")
    for value, weight in weighted:
        if weight <= 0 or not 0 <= value <= 100:
            raise MatchingCalculationError("Une composante du score est hors limites.")
    score = round(sum(value * weight for value, weight in weighted) / total_weight, 2)
    if not 0 <= score <= 100:  # Defensive invariant for future scoring components.
        raise MatchingCalculationError("Le score final calculé est hors limites.")
    return score, total_weight


def calculate_match(application: Application, offer: Offer, analysis: CVAnalysis) -> dict:
    candidate = _skill_set(analysis.skills)
    required = _skill_set(_offer_skills(offer))
    matched_keys = sorted(required.keys() & candidate.keys())
    missing_keys = sorted(required.keys() - candidate.keys())
    additional_keys = sorted(candidate.keys() - required.keys())
    matched = [required[key] for key in matched_keys]
    missing = [required[key] for key in missing_keys]
    additional = [candidate[key] for key in additional_keys]

    components = {}
    weighted = []
    if required:
        skill_score = round(100 * len(matched) / len(required), 2)
        components["skills"] = {"score": skill_score, "weight": 50, "available": True}
        weighted.append((skill_score, 50))
    else:
        components["skills"] = {"score": None, "weight": 50, "available": False}

    years = approximate_experience_years(analysis.experiences)
    required_years = float(offer.required_experience_years) if offer.required_experience_years is not None else None
    experience_score = None
    if required_years is not None and years is not None:
        experience_score = 100.0 if required_years <= 0 or years >= required_years else round(100 * years / required_years, 2)
        weighted.append((experience_score, 25))
    components["experience"] = {
        "score": experience_score, "weight": 25, "available": experience_score is not None,
        "candidate_years": years, "required_years": required_years,
        "explanation": (
            "Comparaison déterministe des années d’expérience datées."
            if required_years is not None
            else "L’offre ne contient pas de critère d’expérience chiffré."
        ),
    }

    required_rank = EDUCATION_RANK.get(offer.required_level)
    candidate_rank = _candidate_education_rank(application, analysis)
    if required_rank:
        education_score = None if candidate_rank is None else (100 if candidate_rank >= required_rank else round(100 * candidate_rank / required_rank, 2))
        components["education"] = {"score": education_score, "weight": 15, "available": education_score is not None}
        if education_score is not None:
            weighted.append((education_score, 15))
    else:
        education_score = None
        components["education"] = {"score": None, "weight": 15, "available": False}
    components["additional"] = {"score": None, "weight": 10, "available": False}
    candidate_text = build_candidate_representation(application, analysis)
    offer_text = build_offer_representation(offer)
    semantic_model = embedding_model_identifier()
    semantic_value = None
    semantic_error = ""
    try:
        semantic_value = calculate_semantic_score(candidate_text, offer_text)
        weighted.append((semantic_value, 30))
    except (EmbeddingError, ValueError) as exc:
        semantic_error = getattr(exc, "code", "embedding_configuration_error")
        logger.warning(
            "Semantic matching unavailable application_id=%s offer_id=%s code=%s detail=%s",
            application.pk,
            offer.pk,
            semantic_error,
            str(exc),
        )
    components["semantic"] = {
        "score": semantic_value,
        "weight": 30,
        "available": semantic_value is not None,
        "model": semantic_model,
        "error_code": semantic_error,
        "explanation": (
            "Similarité cosinus calculée à partir d’embeddings locaux."
            if semantic_value is not None
            else "Embeddings indisponibles ; composante exclue du score."
        ),
    }

    score, total_weight = normalized_weighted_score(weighted)
    components["normalized_weight"] = total_weight
    components["label"] = _score_label(score)
    return {
        "score": score, "matched_skills": matched, "missing_skills": missing,
        "additional_skills": additional, "score_breakdown": components,
        "candidate_summary": _candidate_summary(
            matched=matched, missing=missing, experience_years=years,
            education_score=education_score, score=score,
        ),
        "explanation": (
            f"Score hybride calculé sur {total_weight} point(s) de critères disponibles "
            "(compétences, expérience, formation et similarité sémantique lorsqu’elle est "
            "disponible), puis normalisé sur 100."
        ),
        "algorithm_version": MATCHING_VERSION,
        "candidate_fingerprint": candidate_matching_fingerprint(application, analysis),
        "offer_fingerprint": offer_matching_fingerprint(offer),
        "candidate_representation_hash": text_hash(candidate_text),
        "offer_representation_hash": text_hash(offer_text),
        "semantic_score": semantic_value,
        "semantic_model": semantic_model,
    }


def match_is_stale(match: ApplicationMatch, application: Application | None = None, offer: Offer | None = None, analysis: CVAnalysis | None = None) -> bool:
    """Return whether persisted matching inputs or algorithm have changed."""
    application = application or match.application
    offer = offer or match.offer
    analysis = analysis or CVAnalysis.objects.filter(application=application).first()
    from .analysis_pipeline import analysis_is_stale
    if (
        analysis is None
        or analysis_is_stale(application, analysis)
        or match.algorithm_version != MATCHING_VERSION
        or match.semantic_model != embedding_model_identifier()
    ):
        return True
    return (
        match.candidate_fingerprint != candidate_matching_fingerprint(application, analysis)
        or match.offer_fingerprint != offer_matching_fingerprint(offer)
    )


@transaction.atomic
def match_application(
    application: Application,
    *,
    include_recommendations=True,
    target_offer: Offer | None = None,
) -> list[ApplicationMatch]:
    # Reverse one-to-one relations can keep a stale in-memory instance after a
    # CV review. Matching must always use the latest persisted structured data.
    analysis = CVAnalysis.objects.filter(application=application).first() or extract_cv(application)
    results = []
    offers = Offer.objects.filter(pk=target_offer.pk) if target_offer is not None else Offer.objects.all()
    if target_offer is None and application.offer_id:
        offers = offers.filter(pk=application.offer_id)
    for offer in offers:
        defaults = calculate_match(application, offer, analysis)
        item, _ = ApplicationMatch.objects.update_or_create(
            application=application, offer=offer, defaults=defaults,
        )
        results.append(item)
    if not include_recommendations:
        return results
    update_training_recommendations(application, results)
    return results


def update_training_recommendations(application, results):
    missing_all = {skill for item in results for skill in item.missing_skills}
    from apps.trainings.models import Training
    relevant_ids = []
    for training in Training.objects.all():
        searchable = f"{training.title} {training.description} {training.category} {training.objectives}".casefold()
        covered = sorted(skill for skill in missing_all if skill.casefold() in searchable)
        if covered:
            relevant_ids.append(training.pk)
            TrainingRecommendation.objects.update_or_create(application=application, training=training, defaults={"score": round(100 * len(covered) / len(missing_all), 2), "skill_gaps": covered, "explanation": f"Cette formation couvre {len(covered)} écart(s) de compétences détecté(s)."})
    # Keep human-reviewed history, but remove obsolete pending suggestions.
    TrainingRecommendation.objects.filter(application=application, human_decision="PENDING").exclude(training_id__in=relevant_ids).delete()


def _offer_candidate_pool(offer: Offer) -> list[dict]:
    """Select one stable representative application per candidate profile."""
    applications = Application.objects.select_related(
        "candidate_profile__user", "cv_analysis", "offer"
    ).prefetch_related(
        Prefetch(
            "documents",
            queryset=ApplicationDocument.objects.filter(
                document_type=ApplicationDocumentType.CV,
            ).order_by("-uploaded_at", "-id"),
            to_attr="ranking_cv_documents",
        ),
        Prefetch(
            "matches",
            queryset=ApplicationMatch.objects.filter(offer=offer),
            to_attr="ranking_offer_matches",
        ),
    ).order_by("candidate_profile_id", "-submitted_at", "-id")
    grouped = {}
    for application in applications:
        grouped.setdefault(application.candidate_profile_id, []).append(application)

    pool = []
    for candidate_applications in grouped.values():
        direct_application = next(
            (item for item in candidate_applications if item.offer_id == offer.pk),
            None,
        )
        if direct_application is not None:
            relationship = "APPLIED_TO_OFFER"
            relationship_label = "A postulé à cette offre"
        elif any(item.offer_id is not None for item in candidate_applications):
            relationship = "OTHER_APPLICATION"
            relationship_label = "Autre candidature"
        else:
            relationship = "TALENT_POOL"
            relationship_label = "Vivier de candidats"

        def usable(item):
            analysis = getattr(item, "cv_analysis", None)
            documents = item.ranking_cv_documents
            if analysis is None or not documents:
                return False
            try:
                return (
                    analysis.extractor_version == EXTRACTOR_VERSION
                    and analysis.source_sha256 == cv_document_sha256(documents[0])
                )
            except (OSError, ValueError):
                return False

        usable_applications = [item for item in candidate_applications if usable(item)]
        with_documents = [item for item in candidate_applications if item.ranking_cv_documents]
        representative = next(
            (item for item in usable_applications if item.offer_id == offer.pk),
            usable_applications[0] if usable_applications else next(
                (item for item in with_documents if item.offer_id == offer.pk),
                with_documents[0] if with_documents else direct_application or candidate_applications[0],
            ),
        )
        pool.append({
            "application": representative,
            "relationship": relationship,
            "relationship_label": relationship_label,
            "applied_to_current_offer": direct_application is not None,
        })
    return pool


def rank_offer_candidates(offer: Offer, *, force=False) -> list[dict]:
    """Match the current offer against every usable candidate in the talent pool."""
    from .analysis_pipeline import analysis_is_stale, process_application_analysis
    rows = []
    for pool_item in _offer_candidate_pool(offer):
        application = pool_item["application"]
        if force:
            result = process_application_analysis(application)
            if result.error:
                rows.append({**pool_item, "match": None, "analysis_error": result.error})
                continue
            application.refresh_from_db()
        cv_documents = application.ranking_cv_documents
        analysis = getattr(application, "cv_analysis", None)
        match = next(iter(application.ranking_offer_matches), None)
        if not cv_documents:
            match = None
            analysis_error = "Aucun CV n’est associé à cette candidature."
        elif analysis is None:
            match = None
            analysis_error = "Le CV n’a pas encore pu être analysé."
        elif analysis_is_stale(application, analysis):
            match = None
            analysis_error = "Le CV actuel doit être analysé ou réanalysé."
        else:
            if force or match is None or match_is_stale(match, application=application, offer=offer, analysis=analysis):
                try:
                    match = match_application(
                        application,
                        include_recommendations=False,
                        target_offer=offer,
                    )[0]
                except MatchingCalculationError as exc:
                    match = None
                    analysis_error = str(exc)
                else:
                    analysis_error = ""
            else:
                analysis_error = ""
        rows.append({**pool_item, "match": match, "analysis_error": analysis_error})
    return sorted(
        rows,
        key=lambda row: (
            -(float(row["match"].score) if row["match"] else -1),
            row["application"].submitted_at,
            row["application"].id,
        ),
    )
