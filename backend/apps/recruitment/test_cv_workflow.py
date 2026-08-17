import shutil
import tempfile
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document
import pymupdf
from rest_framework.test import APITestCase

from apps.accounts.choices import UserRole
from .choices import ApplicationDocumentType, ApplicationType
from .intelligence import CVExtractionError, extract_cv_data, parse_cv_text
from .models import Application, ApplicationDocument, CandidateProfile

User = get_user_model()
MEDIA_ROOT = tempfile.mkdtemp()


def docx_cv(name="Jane Doe", email="jane@example.com"):
    stream = BytesIO(); document = Document()
    for line in [name, email, "+212 600 000 000", "Location: Casablanca", "Skills: Python, Django, PostgreSQL",
                 "Experience", "Backend Developer | Acme | 2022 - present", "Education", "Master Informatique | ENSA",
                 "Languages: Français, English", "Certifications: AWS Practitioner"]:
        document.add_paragraph(line)
    document.save(stream)
    return stream.getvalue()


def text_pdf(text: str):
    stream = f"BT /F1 11 Tf 40 760 Td ({text.replace('(', '').replace(')', '')}) Tj ET".encode("latin-1", "replace")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", f"<< /Length {len(stream)} >> stream\n".encode()+stream+b"\nendstream"]
    pdf=bytearray(b"%PDF-1.4\n"); offsets=[]
    for i,obj in enumerate(objects,1): offsets.append(len(pdf));pdf.extend(f"{i} 0 obj\n".encode()+obj+b"\nendobj\n")
    xref=len(pdf);pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets: pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode());return bytes(pdf)


def blank_pdf():
    document = pymupdf.open()
    document.new_page(width=300, height=300)
    content = document.tobytes()
    document.close()
    return content


def two_column_pdf():
    document = pymupdf.open()
    page = document.new_page(width=600, height=800)
    page.insert_textbox(
        (40, 40, 280, 350),
        "COMPÉTENCES\nPython\nDjango\nDISPONIBILITÉS\nWeek-end\nPrésentiel\n"
        "CENTRES D’INTÉRÊT\nLecture\nVoyages",
        fontsize=11,
    )
    page.insert_textbox(
        (320, 40, 570, 500),
        "HIBA AYATALLAH\nhiba.ayatallah@example.com\n06 53 85 49 66\n"
        "EXPÉRIENCES PROFESSIONNELLES\nVENDEUSE\nFASHION NOVA\n2024 - 2025\n"
        "FORMATION\n1ÈRE ANNÉE EN GESTION\nUNIVERSITÉ HASSAN II - CASABLANCA\n"
        "LANGUES\nFrançais : Avancé\nAnglais : Intermédiaire",
        fontsize=11,
    )
    content = document.tobytes()
    document.close()
    return content


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class CVExtractionFormatTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass(); shutil.rmtree(MEDIA_ROOT, ignore_errors=True)

    def test_extracts_docx_and_missing_values_remain_empty(self):
        upload = SimpleUploadedFile("cv.docx", docx_cv(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        data = extract_cv_data(upload)
        self.assertEqual(data["first_name"], "Jane")
        self.assertEqual(data["email"], "jane@example.com")
        self.assertIn("Python", data["skills"])
        self.assertEqual(data["experiences"][0]["company"], "Acme")

        missing = parse_cv_text("Skills: Python")
        self.assertEqual(missing["email"], "")
        self.assertEqual(missing["phone"], "")
        self.assertEqual(missing["certifications"], [])

    def test_extracts_docx_sidebar_table_in_document_order(self):
        stream = BytesIO(); document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "COMPÉTENCES\nPython\nRelation client"
        table.cell(0, 1).text = "HIBA AYATALLAH\nhiba.ayatallah@example.com\n07 12 34 56 78"
        table.cell(1, 0).text = "FORMATION\nLicence en gestion\nUNIVERSITÉ HASSAN II - CASABLANCA"
        table.cell(1, 1).text = "EXPÉRIENCE\nVENDEUSE\nFASHION NOVA\n2024 - 2025"
        document.save(stream)

        upload = SimpleUploadedFile(
            "sidebar.docx", stream.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        data = extract_cv_data(upload)

        self.assertEqual(data["full_name"], "HIBA AYATALLAH")
        self.assertEqual(data["phone"], "+212712345678")
        self.assertTrue(any(item["institution"].casefold() == "université hassan ii" for item in data["education"]))

    def test_extracts_text_pdf_with_pymupdf(self):
        body = "Jane Doe " + "jane@example.com Skills: Python, Django " * 8
        upload = SimpleUploadedFile("cv.pdf", text_pdf(body), content_type="application/pdf")
        data = extract_cv_data(upload)
        self.assertEqual(data["extraction_method"], "PDF_TEXT")
        self.assertEqual(data["email"], "jane@example.com")

    def test_repairs_character_spaced_pdf_text_without_inventing_values(self):
        text = """H I B A
A Y A T A L L A H
a y a t a l l a h . h i b a @ g m a i l . c o m
+ 2 1 2  6  5 3  8 5  4 9  6 6
C a s a b l a n c a  -  M a r o c
C O M P É T E N C E S
L a n g a g e s  :  P y t h o n ,  J a v a ,  C + +
F r a m e w o r k s  :  D j a n g o ,  R e a c t . j s
L A N G U E S
F r a n ç a i s  :  A v a n c é
A n g l a i s  :  I n t e r m é d i a i r e"""
        data = parse_cv_text(text)
        self.assertEqual(data["first_name"], "HIBA")
        self.assertEqual(data["last_name"], "AYATALLAH")
        self.assertEqual(data["email"], "ayatallah.hiba@gmail.com")
        self.assertEqual(data["location"], "Casablanca - Maroc")
        self.assertEqual(data["skills"], ["Python", "Java", "C++", "Django", "React.js"])
        self.assertEqual(data["languages"], ["Français : Avancé", "Anglais : Intermédiaire"])
        self.assertEqual(data["certifications"], [])

    def test_two_column_order_does_not_confuse_experience_education_and_identity(self):
        text = """EXPÉRIENCES PROFESSIONNELLES
VENDEUSE
FASHION NOVA
2024 - 2025
FORMATION
Licence en gestion
UNIVERSITÉ HASSAN II - CASABLANCA
2022 - 2025
HIBA AYATALLAH
hiba.ayatallah@example.com
06 53 85 49 66
COMPÉTENCES
Vente, Relation client"""

        data = parse_cv_text(text)

        self.assertEqual(data["first_name"], "HIBA")
        self.assertEqual(data["last_name"], "AYATALLAH")
        self.assertNotEqual(data["first_name"], "VENDEUSE")
        self.assertNotIn("FASHION NOVA", data["full_name"])
        self.assertEqual(data["phone"], "+212653854966")
        self.assertNotEqual(data["phone"], "20242025")
        self.assertEqual(data["location"].casefold(), "casablanca")
        self.assertTrue(any(item["institution"].casefold() == "université hassan ii" for item in data["education"]))
        self.assertTrue(any(
            item["position"] == "VENDEUSE"
            and item["company"] == "FASHION NOVA"
            and item["start_date"] == "2024"
            and item["end_date"] == "2025"
            for item in data["experiences"]
        ))

    def test_phone_validation_normalizes_moroccan_formats_and_rejects_years(self):
        for raw, expected in [
            ("+212 7 12 34 56 78", "+212712345678"),
            ("07-12-34-56-78", "+212712345678"),
            ("06 12 34 56 78", "+212612345678"),
        ]:
            with self.subTest(raw=raw):
                self.assertEqual(parse_cv_text(f"Jane Doe\njane@example.com\n{raw}")["phone"], expected)
        self.assertEqual(parse_cv_text("Jane Doe\njane@example.com\n2024 - 2025\n20242025")["phone"], "")

    def test_uncertain_identity_and_institution_address_are_left_safe(self):
        data = parse_cv_text("""VENDEUSE
FASHION NOVA
2024 - 2025
FORMATION
UNIVERSITÉ HASSAN II - CASABLANCA
contact@example.com""")

        self.assertEqual(data["full_name"], "")
        self.assertEqual(data["phone"], "")
        self.assertEqual(data["location"].casefold(), "casablanca")
        self.assertNotIn("université", data["location"].casefold())

    def test_pdf_without_exploitable_text_returns_clear_message(self):
        upload = SimpleUploadedFile("scan.pdf", blank_pdf(), content_type="application/pdf")
        with self.assertRaisesRegex(
            CVExtractionError,
            "Ce PDF ne contient pas de texte exploitable",
        ):
            extract_cv_data(upload)

    def test_sections_stop_at_interests_availability_and_new_headings(self):
        data = parse_cv_text("""HIBA AYATALLAH
hiba.ayatallah@example.com
COMPÉTENCES
Python | Django | | |
DISPONIBILITÉS
Week-end
Présentiel
Horaires décalés
Travail le soir
LANGUES
Français : Avancé
Anglais : Intermédiaire
CENTRES D’INTÉRÊT
Lecture
FORMATION
1ÈRE ANNÉE EN GESTION
UNIVERSITÉ HASSAN II - CASABLANCA""")

        self.assertEqual(data["skills"], ["Python", "Django"])
        self.assertEqual(data["languages"], ["Français : Avancé", "Anglais : Intermédiaire"])
        self.assertEqual(len(data["education"]), 1)
        self.assertEqual(data["education"][0]["title"], "1ÈRE ANNÉE EN GESTION")
        self.assertEqual(data["education"][0]["institution"], "UNIVERSITÉ HASSAN II")

    def test_experiences_are_grouped_cleaned_and_deduplicated(self):
        data = parse_cv_text("""Jane Doe
jane@example.com
EXPÉRIENCES
Développeuse Backend | Acme | | | 2022 - 2024
Développeuse Backend
Acme
2022 - 2024
Consultante
Beta SARL
2024 - présent
FORMATION
Master Informatique
ENSA""")

        self.assertEqual(len(data["experiences"]), 2)
        self.assertNotIn("|", "\n".join(item["description"] for item in data["experiences"]))
        self.assertEqual(data["experiences"][0]["company"], "Acme")
        self.assertEqual(data["experiences"][1]["company"], "Beta SARL")

    def test_two_column_pdf_uses_blocks_and_keeps_sections_isolated(self):
        upload = SimpleUploadedFile("columns.pdf", two_column_pdf(), content_type="application/pdf")
        data = extract_cv_data(upload)

        self.assertEqual(data["full_name"], "HIBA AYATALLAH")
        self.assertEqual(data["skills"], ["Python", "Django"])
        self.assertEqual(data["languages"], ["Français : Avancé", "Anglais : Intermédiaire"])
        self.assertEqual(len(data["experiences"]), 1)
        self.assertEqual(data["experiences"][0]["company"], "FASHION NOVA")
        self.assertEqual(data["education"][0]["title"], "1ÈRE ANNÉE EN GESTION")

    def test_name_is_detected_when_first_and_last_names_are_split(self):
        data = parse_cv_text("""COMPÉTENCES
Python
CONTACT
HIBA
AYATALLAH
hiba.ayatallah@example.com
06 53 85 49 66""")
        self.assertEqual(data["first_name"], "HIBA")
        self.assertEqual(data["last_name"], "AYATALLAH")

    def test_template_one_competition_is_one_experience_and_education_is_semantic(self):
        data = parse_cv_text("""HIBA AYATALLAH
hiba.ayatallah@example.com
PARCOURS ACADÉMIQUE
CYCLE INGÉNIEUR – IADATA 2024 – PRÉSENT
EMSI CASABLANCA
CYCLE PRÉPARATOIRE EN INGÉNIERIE INFORMATIQUE
EMSI CASABLANCA
BACCALAURÉAT SCIENCES PHYSIQUES ET CHIMIQUES
GROUPE SCOLAIRE ANOUAR AL MADINA
EXPÉRIENCE PROFESSIONNELLE
Plateforme de gestion de l’enseignement en ligne
Juillet – Septembre 2025
EMSI Intership Summer Compétition — 1er Prix
Développement sous Moodle
COMPÉTENCES TECHNIQUES
React Native, Android Studio, GitHub""")

        self.assertEqual(len(data["experiences"]), 1)
        self.assertEqual(data["experiences"][0]["position"], "")
        self.assertEqual(data["experiences"][0]["company"], "")
        self.assertEqual(len(data["education"]), 3)
        self.assertIn("React Native", data["skills"])
        self.assertIn("Android Studio", data["skills"])
        self.assertIn("GitHub", data["skills"])

    def test_template_two_headings_and_degrees_are_not_experiences_or_skills(self):
        data = parse_cv_text("""YASMINE CHAKOURI
yasminech37@gmail.com
FORMATION
1ÈRE ANNÉE EN GESTION | UNIVERSITÉ HASSAN II
2025 - Présent
EXPÉRIENCE
VENDEUSE - FASHION NOVA
Casablanca | Mai 2025 - Octobre 2025
BACCALAURÉAT EN SCIENCES
ÉCONOMIQUES - MENTION BIEN
2024 - 2025
COMPÉTENCES PROFESSIONNELLES
Microsoft Word
Microsoft Excel
LANGUES
Arabe : Langue maternelle""")

        self.assertNotIn("COMPÉTENCES PROFESSIONNELLES", data["skills"])
        self.assertEqual(len(data["experiences"]), 1)
        self.assertEqual(data["experiences"][0]["position"], "VENDEUSE")
        self.assertEqual(data["experiences"][0]["company"], "FASHION NOVA")
        self.assertTrue(any("BACCALAURÉAT" in item["title"] for item in data["education"]))
        self.assertTrue(any(item["institution"] == "UNIVERSITÉ HASSAN II" for item in data["education"]))


@override_settings(MEDIA_ROOT=MEDIA_ROOT)
class CVWorkflowAPITests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(email="cv-admin@test.com", password="test")
        self.candidate = User.objects.create_user(email="cv-candidate@test.com", password="test", role=UserRole.CANDIDATE)
        self.other = User.objects.create_user(email="cv-other@test.com", password="test", role=UserRole.CANDIDATE)
        self.profile = CandidateProfile.objects.create(user=self.candidate, phone_number="", current_school="", study_level="OTHER", study_field="")
        other_profile = CandidateProfile.objects.create(user=self.other, phone_number="", current_school="", study_level="OTHER", study_field="")
        self.application = Application.objects.create(candidate_profile=self.profile, application_type=ApplicationType.HIRING)
        self.other_application = Application.objects.create(candidate_profile=other_profile, application_type=ApplicationType.HIRING)

    def upload(self, user, application=None):
        self.client.force_authenticate(user)
        file = SimpleUploadedFile("candidate.docx", docx_cv(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        return self.client.post(f"/api/applications/{(application or self.application).id}/upload-cv/", {"file": file}, format="multipart")

    def test_candidate_uploads_edits_validates_and_persists_profile(self):
        response = self.upload(self.candidate)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["analysis"]["email"], "jane@example.com")
        self.assertEqual(response.data["analysis"]["first_name"], "Jane")
        patch_response = self.client.patch(f"/api/applications/{self.application.id}/cv-analysis/", {"phone": "+212611111111", "skills": ["Python", "Django"]}, format="json")
        self.assertEqual(patch_response.status_code, 200)
        validated = self.client.post(f"/api/applications/{self.application.id}/validate-cv/", {"first_name": "Jane", "last_name": "Doe", "location": "Rabat"}, format="json")
        self.assertEqual(validated.status_code, 200)
        self.profile.refresh_from_db(); self.candidate.refresh_from_db()
        self.assertEqual(self.profile.address, "Rabat")
        self.assertEqual(self.candidate.first_name, "Jane")
        self.assertTrue(self.application.cv_analysis.human_validated)

    def test_candidate_cannot_access_another_candidates_cv(self):
        self.upload(self.admin, self.other_application)
        self.client.force_authenticate(self.candidate)
        self.assertEqual(self.client.get(f"/api/applications/{self.other_application.id}/cv-analysis/").status_code, 404)

    def test_super_admin_can_upload_and_edit_candidate_cv(self):
        response = self.upload(self.admin)
        self.assertEqual(response.status_code, 201)
        response = self.client.patch(f"/api/applications/{self.application.id}/cv-analysis/", {"languages": ["Français", "English"]}, format="json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["languages"], ["Français", "English"])

    def test_invalid_cv_is_rejected(self):
        self.client.force_authenticate(self.candidate)
        bad = SimpleUploadedFile("broken.pdf", b"%PDF broken", content_type="application/pdf")
        response = self.client.post(f"/api/applications/{self.application.id}/upload-cv/", {"file": bad}, format="multipart")
        self.assertEqual(response.status_code, 400)
